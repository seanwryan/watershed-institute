#!/usr/bin/env python3
"""Focused tests for final All StreamWatch Data.xlsx ETL compatibility."""

from __future__ import annotations

import inspect
import sys
import tempfile
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from etl.all_data_workbook import (
    CHLORIDE_APPLY_DIVIDE_BY_TEN,
    FINAL_METHOD_NAMES,
    NITRATE_UNIT_POLICY,
    classify_data_condition,
    dry_run_all_data,
    find_all_data_header_row,
    map_chloride_mg_l,
    mapped_chem_values,
    parse_e_coli_result,
    read_all_data_sheet,
)
from etl.chem_recon import CHEM_HEADER_ALIASES, chem_fingerprint, resolve_data_condition_id
from etl.visit_helpers import insert_bacteria

FINAL_WB = Path(
    "/Users/seanryan/Downloads/Data Work Folder 2/All StreamWatch Data.xlsx"
)
FLAT_CANDIDATE = Path(
    "/Users/seanryan/Downloads/Data Work Folder/All StreamWatch Data.xlsx"
)


def _write_banner_workbook(path: Path):
    """Minimal workbook mimicking final banner + header layout."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "ALL DATA"
    for i in range(1, 15):
        ws.cell(i, 1, f"Banner row {i}")
    headers = [
        "Data Condition",
        "Notes",
        "Method",
        "Site",
        "Date",
        "Air Temperature (°C)",
        "Water Temperature (°C)",
        "Nitrate (mg/L)",
        "Phosphate (mg/L)",
        "pH",
        "Turbidity (JTU/NTU)",
        "DO (ppm)",
        "DO (%)",
        "Conductivity (µS/cm)",
        "Chloride (mg/L)",
        "E. coli Result",
        "E. coli Mod.",
    ]
    for c, h in enumerate(headers, 1):
        ws.cell(15, c, h)
    data_rows = [
        [
            "Provisional",
            None,
            "CAT: Hanna",
            "AC1",
            "2024-06-16",
            20.0,
            18.5,
            0.5,
            0.1,
            7.2,
            3.0,
            8.1,
            90.0,
            400.0,
            49.8565,
            120,
            "=",
        ],
        [
            "Outlier",
            None,
            "CAT: LaMotte",
            "AC1",
            "2024-06-16",
            21.0,
            19.0,
            0.6,
            0.2,
            7.0,
            4.0,
            7.5,
            None,
            None,
            None,
            None,
            None,
        ],
        [
            "Flagged; Incomplete",
            None,
            "Salt Watch",
            "SB2",
            "2024-07-01",
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            43.0,
            None,
            None,
        ],
        [
            "Duplicate?",
            None,
            "BACT",
            "HO2",
            "2025-07-13",
            None,
            22.0,
            0.3,
            0.05,
            None,
            2.0,
            None,
            None,
            None,
            29.6,
            2419,
            ">",
        ],
        [
            "Provisional",
            None,
            "CAT: Hanna",
            "AC1",
            "2024-06-16",
            20.0,
            18.5,
            0.5,
            0.1,
            7.2,
            3.0,
            8.1,
            90.0,
            400.0,
            49.8565,
            None,
            None,
        ],
    ]
    for r, row in enumerate(data_rows, start=16):
        for c, val in enumerate(row, 1):
            cell = ws.cell(r, c, val)
            if isinstance(val, str) and val in {"=", ">", "<"}:
                cell.value = val
                cell.data_type = "s"
    wb.save(path)


def test_header_detection_banner_and_flat():
    with tempfile.TemporaryDirectory() as td:
        p = Path(td) / "banner.xlsx"
        _write_banner_workbook(p)
        assert find_all_data_header_row(p) == 14
        df, hdr = read_all_data_sheet(p)
        assert hdr == 14
        assert "Method" in df.columns
        assert "Air Temperature (°C)" in df.columns
        assert len(df) == 5

    if FLAT_CANDIDATE.exists():
        assert find_all_data_header_row(FLAT_CANDIDATE) == 0


def test_unit_suffixed_aliases():
    with tempfile.TemporaryDirectory() as td:
        p = Path(td) / "banner.xlsx"
        _write_banner_workbook(p)
        df, _ = read_all_data_sheet(p)
        row = df.iloc[0]
        vals = mapped_chem_values(row)
        assert vals["air_temp_c"] == 20.0
        assert vals["water_temp_c"] == 18.5
        assert vals["nitrate_ug_l"] == 0.5  # stored as-is; no *1000
        assert vals["phosphate_mg_l"] == 0.1
        assert vals["ph"] == 7.2
        assert vals["turbidity_ntu"] == 3.0
        assert vals["dissolved_oxygen_ppm"] == 8.1
        assert vals["dissolved_oxygen_pct"] == 90.0
        assert vals["conductivity_us_cm"] == 400.0
        assert vals["chloride_mg_l"] == 49.8565


def test_method_names_and_salt_watch():
    assert "CAT: Early LaMotte" in FINAL_METHOD_NAMES
    assert "CAT: LaMotte" in FINAL_METHOD_NAMES
    assert "CAT: Hanna" in FINAL_METHOD_NAMES
    assert "Salt Watch" in FINAL_METHOD_NAMES
    with tempfile.TemporaryDirectory() as td:
        p = Path(td) / "banner.xlsx"
        _write_banner_workbook(p)
        stats = dry_run_all_data(p)
        assert stats["method_counts"]["CAT: Hanna"] == 2  # incl exact dup row
        assert stats["method_counts"]["Salt Watch"] == 1
        assert stats["method_counts"]["CAT: LaMotte"] == 1


def test_e_coli_result_and_mod():
    with tempfile.TemporaryDirectory() as td:
        p = Path(td) / "banner.xlsx"
        _write_banner_workbook(p)
        df, _ = read_all_data_sheet(p)
        e1, m1, s1 = parse_e_coli_result(df.iloc[0])
        assert e1 == 120 and m1 == "=" and s1 is None
        e2, m2, s2 = parse_e_coli_result(df.iloc[3])
        assert e2 == 2419 and m2 == ">" and s2 is None
        stats = dry_run_all_data(p)
        assert stats["e_coli_result_parseable"] >= 2
        assert stats["e_coli_with_modifier"] >= 2
    # insert_bacteria accepts detection_limit_note
    assert "detection_limit_note" in inspect.signature(insert_bacteria).parameters


def test_data_conditions_outlier_duplicate_and_compounds():
    cond_map = {
        "Provisional": 1,
        "Outlier": 2,
        "Duplicate?": 3,
        "Flagged": 4,
        "Incomplete": 5,
        "Minor_Deviation": 6,
    }
    unresolved = []
    assert classify_data_condition("Outlier", cond_map, unresolved)["status"] == "resolved"
    assert classify_data_condition("Duplicate?", cond_map, unresolved)["status"] == "resolved"
    assert (
        classify_data_condition("Minor Deviation", cond_map, unresolved)["status"]
        == "resolved"
    )
    compound = classify_data_condition("Flagged; Incomplete", cond_map, unresolved)
    assert compound["status"] == "unresolved_compound"
    assert compound["data_condition_id"] is None
    assert "Flagged; Incomplete" in unresolved
    # resolve_data_condition_id must not invent first-token mapping
    unresolved2 = []
    assert resolve_data_condition_id("Flagged; Incomplete", cond_map, unresolved2) is None


def test_chloride_no_double_correction():
    assert CHLORIDE_APPLY_DIVIDE_BY_TEN is False
    assert map_chloride_mg_l(49.8565) == 49.8565
    assert map_chloride_mg_l(498.565) == 498.565  # would be wrong if ÷10 applied
    # Guard: migrate module must not contain a chloride ÷10 transform
    src = Path("etl/all_data_workbook.py").read_text(encoding="utf-8")
    assert "CHLORIDE_APPLY_DIVIDE_BY_TEN = False" in src
    assert "/ 10" not in src.replace("÷10", "")


def test_multipackage_fingerprint_and_dedupe():
    with tempfile.TemporaryDirectory() as td:
        p = Path(td) / "banner.xlsx"
        _write_banner_workbook(p)
        df, _ = read_all_data_sheet(p)
        v0 = mapped_chem_values(df.iloc[0])
        v1 = mapped_chem_values(df.iloc[1])
        v_dup = mapped_chem_values(df.iloc[4])
        fp0 = chem_fingerprint("AC1", pd.Timestamp("2024-06-16").date(), "CAT: Hanna", v0)
        fp1 = chem_fingerprint("AC1", pd.Timestamp("2024-06-16").date(), "CAT: LaMotte", v1)
        fp_dup = chem_fingerprint(
            "AC1", pd.Timestamp("2024-06-16").date(), "CAT: Hanna", v_dup
        )
        assert fp0 != fp1  # legitimate multi-package
        assert fp0 == fp_dup
        stats = dry_run_all_data(p)
        assert stats["exact_duplicates"] == 1
        assert stats["multi_package_site_dates"] >= 1
        assert stats["would_insert_chemistry"] == 4  # Hanna, LaMotte, Salt Watch, BACT


def test_dry_run_no_write_and_nitrate_policy():
    assert "without unit conversion" in NITRATE_UNIT_POLICY.lower() or "as-is" in NITRATE_UNIT_POLICY
    with tempfile.TemporaryDirectory() as td:
        p = Path(td) / "banner.xlsx"
        _write_banner_workbook(p)
        stats = dry_run_all_data(p)
        assert stats["writes_database"] is False
        assert stats["workbook_rows_read"] == 5
        assert stats["chemistry_capable_rows"] >= 3
        assert stats["chloride_nonnull"] >= 2
        assert "Air Temperature (°C)" in CHEM_HEADER_ALIASES["air_temp_c"]
        assert "Nitrate (mg/L)" in CHEM_HEADER_ALIASES["nitrate_ug_l"]


def test_final_workbook_if_present():
    if not FINAL_WB.exists():
        print("SKIP final workbook dry-run (file not present)")
        return
    assert find_all_data_header_row(FINAL_WB) == 14
    stats = dry_run_all_data(FINAL_WB)
    assert stats["writes_database"] is False
    assert stats["header_row_0based"] == 14
    assert stats["workbook_rows_read"] >= 18000
    assert stats["chemistry_capable_rows"] > 15000
    assert stats["parameter_non_null"]["water_temp_c"] > 10000
    assert stats["parameter_non_null"]["chloride_mg_l"] > 2000
    assert stats["e_coli_result_parseable"] > 2000
    assert stats["method_counts"].get("CAT: Early LaMotte", 0) > 8000
    assert stats["method_counts"].get("Salt Watch", 0) > 300
    assert stats["data_condition_status_counts"].get("resolved", 0) > 10000
    assert stats["data_condition_status_counts"].get("unresolved_compound", 0) > 0
    # Outlier / Duplicate? should resolve when default dry-run cond map includes them
    # Compounds remain unresolved
    assert any(";" in s for s in stats["unresolved_data_conditions"])


def main():
    test_header_detection_banner_and_flat()
    test_unit_suffixed_aliases()
    test_method_names_and_salt_watch()
    test_e_coli_result_and_mod()
    test_data_conditions_outlier_duplicate_and_compounds()
    test_chloride_no_double_correction()
    test_multipackage_fingerprint_and_dedupe()
    test_dry_run_no_write_and_nitrate_policy()
    test_final_workbook_if_present()
    print("ALL_DATA_ETL_COMPAT_OK")


if __name__ == "__main__":
    main()
